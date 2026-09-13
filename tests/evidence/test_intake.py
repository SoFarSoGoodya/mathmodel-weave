from datetime import date
import csv
import json

from openpyxl import Workbook
from pypdf import PdfWriter

from mathmodel_agent.evidence import (
    correct_pdf_text,
    intake_pdf,
    intake_xlsx,
    request_visual_review,
    validate_xlsx_export,
)


def test_pdf_extraction_correction_and_targeted_review(tmp_path):
    source = tmp_path / "source.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    with source.open("wb") as stream:
        writer.write(stream)

    case = tmp_path / "case"
    record = intake_pdf(case, source, problem_id="question")
    assert (case / record["raw"]).read_bytes() == source.read_bytes()
    assert "<!-- page: 1 -->" in (case / record["extracted"]).read_text()
    corrected = correct_pdf_text(case, "question", "<!-- page: 1 -->\n\nCorrected problem text")
    assert "Corrected" in (case / corrected["corrected"]).read_text()
    assert "Corrected" in (case / corrected["diff"]).read_text()
    request = request_visual_review(case, "question", 1, "table at top", "Read the missing unit.")
    assert request["mode"] == "targeted_visual_review"
    assert (case / "problem" / "review_requests" / f"{request['review_id']}.json").exists()


def test_xlsx_streams_hidden_sheet_dates_formulas_and_detects_truncation(tmp_path):
    source = tmp_path / "input.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Measurements"
    sheet.append(["id", "value", "observed", "formula"])
    for index in range(1, 251):
        sheet.append([index, index * 1.25, date(2026, 1, 1), f"=B{index + 1}*2"])
    hidden = workbook.create_sheet("Reference")
    hidden.sheet_state = "hidden"
    hidden.append(["constant", 42])
    workbook.create_sheet("Empty")
    workbook.save(source)

    case = tmp_path / "case"
    record = intake_xlsx(case, source, dataset_id="data", preview_rows=2, preview_columns=3)
    sheets = {item["name"]: item for item in record["sheets"]}
    measured, hidden_sheet, empty = sheets["Measurements"], sheets["Reference"], sheets["Empty"]
    assert measured["source_nonempty_cells"] == 1004
    assert measured["formula_cells"] == measured["formula_without_cache"] == 250
    assert measured["type_counts"]["datetime"] == 250
    assert measured["schema"][2]["type_counts"]["datetime"] == 250
    assert hidden_sheet["hidden"] is True
    assert hidden_sheet["source_nonempty_cells"] == 2
    assert empty["bounds"] is None
    assert record["verification"]["ok"] is True
    registry = json.loads((case / "knowledge" / "registry.json").read_text())
    assert registry["datasets"]["data"]["current_version"] == "v1"

    csv_path = case / measured["csv"]
    rows = list(csv.reader(csv_path.open(encoding="utf-8", newline="")))
    assert rows[1][1] == "1.25"
    assert rows[1][2] == "2026-01-01T00:00:00"
    assert rows[1][3] == "=B2*2"
    with csv_path.open("w", encoding="utf-8", newline="") as stream:
        csv.writer(stream).writerows(rows[:-1])
    assert validate_xlsx_export(case, "data")["ok"] is False
