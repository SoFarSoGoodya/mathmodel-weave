"""Faithful PDF and XLSX intake with bounded profiles and conversion checks."""

from __future__ import annotations

import csv
from datetime import date, datetime, time
import difflib
import hashlib
from itertools import zip_longest
from pathlib import Path
import shutil
from typing import Any

from openpyxl import load_workbook
from pypdf import PdfReader

from mathmodel_agent.contracts import read_json, sha256_file, write_json

from ._store import case_path, relative_to_case, slug
from .pool import register_dataset


def _copy_original(root: Path, source: str | Path, destination: Path) -> Path:
    source = Path(source).expanduser().resolve()
    if not source.is_file():
        raise FileNotFoundError(source)
    destination.parent.mkdir(parents=True, exist_ok=True)
    if destination.exists() and sha256_file(destination) != sha256_file(source):
        raise FileExistsError(f"refusing to overwrite preserved original: {destination}")
    if not destination.exists():
        shutil.copy2(source, destination)
    return destination


def intake_pdf(case_root: str | Path, pdf_path: str | Path, *, problem_id: str | None = None) -> dict:
    """Copy a PDF unchanged and create a simple page-marked Markdown extraction."""
    root = case_path(case_root)
    source = Path(pdf_path).expanduser().resolve()
    problem_id = problem_id or slug(source.stem, fallback="problem")
    raw = _copy_original(root, source, root / "problem" / "raw" / f"{problem_id}{source.suffix.lower()}")
    reader = PdfReader(raw)
    pages = []
    for number, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        pages.append(f"<!-- page: {number} -->\n\n{text.rstrip()}")
    extracted = root / "problem" / f"{problem_id}.extracted.md"
    extracted.parent.mkdir(parents=True, exist_ok=True)
    extracted.write_text("\n\n".join(pages).rstrip() + "\n", encoding="utf-8")
    manifest = {
        "problem_id": problem_id,
        "raw": relative_to_case(root, raw),
        "raw_sha256": sha256_file(raw),
        "extracted": relative_to_case(root, extracted),
        "page_count": len(reader.pages),
        "status": "needs_review",
        "limitations": ["pypdf extracts embedded text only; it does not OCR scanned pages."],
    }
    write_json(root / "problem" / f"{problem_id}.json", manifest)
    return manifest


def correct_pdf_text(
    case_root: str | Path,
    problem_id: str,
    corrected_markdown: str | Path,
    *,
    reviewer: str | None = None,
) -> dict:
    """Store a reviewed Markdown version and a unified diff from extracted text."""
    root = case_path(case_root)
    manifest_path = root / "problem" / f"{problem_id}.json"
    manifest = read_json(manifest_path)
    candidate = Path(corrected_markdown)
    corrected = candidate.read_text(encoding="utf-8") if candidate.exists() else str(corrected_markdown)
    extracted = (root / manifest["extracted"]).read_text(encoding="utf-8")
    corrected_path = root / "problem" / f"{problem_id}.corrected.md"
    diff_path = root / "problem" / f"{problem_id}.corrections.diff"
    corrected_path.write_text(corrected.rstrip() + "\n", encoding="utf-8")
    diff = difflib.unified_diff(
        extracted.splitlines(keepends=True), corrected_path.read_text(encoding="utf-8").splitlines(keepends=True),
        fromfile=manifest["extracted"], tofile=relative_to_case(root, corrected_path),
    )
    diff_path.write_text("".join(diff), encoding="utf-8")
    manifest.update(
        {
            "corrected": relative_to_case(root, corrected_path),
            "diff": relative_to_case(root, diff_path),
            "reviewer": reviewer,
            "status": "corrected",
        }
    )
    write_json(manifest_path, manifest)
    return manifest


def request_visual_review(
    case_root: str | Path,
    problem_id: str,
    page: int,
    region: str | dict,
    question: str,
) -> dict:
    """Record one explicitly bounded visual-review request; it never requests all pages."""
    if page < 1 or not question.strip():
        raise ValueError("page must be positive and question is required")
    root = case_path(case_root)
    manifest = read_json(root / "problem" / f"{problem_id}.json")
    if page > manifest["page_count"]:
        raise ValueError("page is outside this PDF")
    review_id = f"review-{hashlib.sha256(f'{problem_id}:{page}:{region}:{question}'.encode()).hexdigest()[:12]}"
    request = {"review_id": review_id, "problem_id": problem_id, "page": page, "region": region, "question": question, "mode": "targeted_visual_review"}
    path = root / "problem" / "review_requests" / f"{review_id}.json"
    write_json(path, request)
    return request


def _type_name(value: Any) -> str:
    if value is None:
        return "blank"
    if isinstance(value, bool):
        return "bool"
    if isinstance(value, datetime):
        return "datetime"
    if isinstance(value, date):
        return "date"
    if isinstance(value, time):
        return "time"
    if isinstance(value, int):
        return "int"
    if isinstance(value, float):
        return "float"
    return "text"


def _csv_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, bool):
        return "true" if value else "false"
    return str(value)


def _sheet_name(name: str, index: int) -> str:
    return f"{index + 1:02d}-{slug(name, fallback='sheet')}"


def _rows_for_sheet(ws_formula, ws_cached):
    for formula_row, cached_row in zip_longest(ws_formula.iter_rows(), ws_cached.iter_rows(), fillvalue=()):
        if len(formula_row) != len(cached_row):
            raise ValueError("formula and cached workbook views have different dimensions")
        yield formula_row, cached_row


def _stream_sheet(ws_formula, ws_cached, csv_path: Path, preview_rows: int, preview_columns: int) -> dict:
    min_row = min_col = None
    max_row = max_col = 0
    nonempty = 0
    formulas = missing_cache = 0
    type_counts: dict[str, int] = {}
    column_type_counts: dict[int, dict[str, int]] = {}
    csv_path.parent.mkdir(parents=True, exist_ok=True)
    temporary = csv_path.with_suffix(".all.tmp")
    with temporary.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.writer(stream)
        for row_index, (formula_row, cached_row) in enumerate(_rows_for_sheet(ws_formula, ws_cached), start=1):
            output_row: list[str] = []
            for column_index, (formula_cell, cached_cell) in enumerate(zip(formula_row, cached_row), start=1):
                formula_value = formula_cell.value
                is_formula = formula_cell.data_type == "f"
                if is_formula:
                    formulas += 1
                    if cached_cell.value is None:
                        missing_cache += 1
                        value, kind = formula_value, "formula_without_cache"
                    else:
                        value, kind = cached_cell.value, "formula_cached"
                else:
                    value, kind = formula_value, _type_name(formula_value)
                output_row.append(_csv_value(value))
                type_counts[kind] = type_counts.get(kind, 0) + 1
                column_counts = column_type_counts.setdefault(column_index, {})
                column_counts[kind] = column_counts.get(kind, 0) + 1
                if value is not None:
                    nonempty += 1
                    min_row = row_index if min_row is None else min(min_row, row_index)
                    min_col = column_index if min_col is None else min(min_col, column_index)
                    max_row, max_col = max(max_row, row_index), max(max_col, column_index)
            writer.writerow(output_row)
    preview: list[list[str]] = []
    if min_row is None:
        bounds = None
        exported_rows = exported_columns = 0
        csv_path.write_text("", encoding="utf-8")
    else:
        bounds = {"min_row": min_row, "max_row": max_row, "min_column": min_col, "max_column": max_col}
        exported_rows, exported_columns = max_row - min_row + 1, max_col - min_col + 1
        with temporary.open(encoding="utf-8", newline="") as input_stream, csv_path.open("w", encoding="utf-8", newline="") as output_stream:
            reader, writer = csv.reader(input_stream), csv.writer(output_stream)
            for row_index, row in enumerate(reader, start=1):
                if min_row <= row_index <= max_row:
                    exported = row[min_col - 1:max_col]
                    writer.writerow(exported)
                    if len(preview) < preview_rows:
                        preview.append([value[:160] for value in exported[:preview_columns]])
    temporary.unlink()
    return {
        "bounds": bounds,
        "source_nonempty_cells": nonempty,
        "type_counts": type_counts,
        "schema": [
            {"column": column, "type_counts": counts}
            for column, counts in sorted(column_type_counts.items())
        ],
        "formula_cells": formulas,
        "formula_without_cache": missing_cache,
        "preview": preview,
        "csv_rows": exported_rows,
        "csv_columns": exported_columns,
    }


def intake_xlsx(
    case_root: str | Path,
    xlsx_path: str | Path,
    *,
    dataset_id: str | None = None,
    preview_rows: int = 5,
    preview_columns: int = 8,
) -> dict:
    """Stream every visible and hidden sheet to CSV without cleaning or recalculating formulas."""
    if preview_rows < 1 or preview_columns < 1:
        raise ValueError("preview sizes must be positive")
    root = case_path(case_root)
    source = Path(xlsx_path).expanduser().resolve()
    dataset_id = dataset_id or slug(source.stem, fallback="dataset")
    raw = _copy_original(root, source, root / "datasets" / "raw" / f"{dataset_id}{source.suffix.lower()}")
    formulas = load_workbook(raw, read_only=True, data_only=False)
    cached = load_workbook(raw, read_only=True, data_only=True)
    sheets = []
    derived = []
    try:
        for index, name in enumerate(formulas.sheetnames):
            output = root / "datasets" / "derived" / dataset_id / "csv" / f"{_sheet_name(name, index)}.csv"
            profile = _stream_sheet(formulas[name], cached[name], output, preview_rows, preview_columns)
            profile.update({"name": name, "sheet_key": _sheet_name(name, index), "hidden": formulas[name].sheet_state != "visible", "csv": relative_to_case(root, output)})
            sheets.append(profile)
            derived.append(output)
    finally:
        formulas.close()
        cached.close()
    profile_path = root / "datasets" / "derived" / dataset_id / "profile.json"
    manifest = {
        "dataset_id": dataset_id,
        "raw": relative_to_case(root, raw),
        "raw_sha256": sha256_file(raw),
        "format": "xlsx",
        "conversion": "csv_per_sheet",
        "sheets": sheets,
        "limitations": ["openpyxl reads formula text and cached values but does not calculate formulas."],
    }
    write_json(profile_path, manifest)
    verification = validate_xlsx_export(root, dataset_id, profile=manifest)
    manifest["verification"] = verification
    write_json(profile_path, manifest)
    register_dataset(root, dataset_id, raw_path=raw, profile_path=profile_path, derived_paths=derived)
    return manifest


def validate_xlsx_export(case_root: str | Path, dataset_id: str, *, profile: dict | None = None) -> dict:
    """Independently rescan source workbook and CSV values to catch dropped cells or regions."""
    root = case_path(case_root)
    profile = profile or read_json(root / "datasets" / "derived" / dataset_id / "profile.json")
    formulas = load_workbook(root / profile["raw"], read_only=True, data_only=False)
    cached = load_workbook(root / profile["raw"], read_only=True, data_only=True)
    checks = []
    try:
        for expected in profile["sheets"]:
            ws_formula, ws_cached = formulas[expected["name"]], cached[expected["name"]]
            bounds = expected["bounds"]
            if bounds is None:
                with (root / expected["csv"]).open(encoding="utf-8", newline="") as stream:
                    ok = next(csv.reader(stream), None) is None
                checks.append({"sheet": expected["name"], "ok": ok, "reason": "empty sheet" if ok else "CSV has data for empty source"})
                continue
            height, width = bounds["max_row"] - bounds["min_row"] + 1, bounds["max_column"] - bounds["min_column"] + 1
            expected_nonempty = 0
            actual_type_counts: dict[str, int] = {}
            actual_count, ok = 0, True
            with (root / expected["csv"]).open(encoding="utf-8", newline="") as stream:
                actual_reader = csv.reader(stream)
                for row_index, (formula_row, cached_row) in enumerate(_rows_for_sheet(ws_formula, ws_cached), start=1):
                    if not bounds["min_row"] <= row_index <= bounds["max_row"]:
                        continue
                    actual_row = next(actual_reader, None)
                    actual_count += actual_row is not None
                    ok = ok and actual_row is not None and len(actual_row) == width
                    for column_index in range(bounds["min_column"], bounds["max_column"] + 1):
                        formula_cell, cached_cell = formula_row[column_index - 1], cached_row[column_index - 1]
                        value = cached_cell.value if formula_cell.data_type == "f" and cached_cell.value is not None else formula_cell.value
                        kind = "formula_cached" if formula_cell.data_type == "f" and cached_cell.value is not None else ("formula_without_cache" if formula_cell.data_type == "f" else _type_name(formula_cell.value))
                        actual_type_counts[kind] = actual_type_counts.get(kind, 0) + 1
                        expected_nonempty += value is not None
                        actual_text = actual_row[column_index - bounds["min_column"]] if actual_row and len(actual_row) == width else None
                        ok = ok and actual_text == _csv_value(value)
                ok = ok and next(actual_reader, None) is None and actual_count == height
            ok = ok and actual_type_counts == expected["type_counts"]
            checks.append({"sheet": expected["name"], "ok": bool(ok), "source_nonempty_cells": expected_nonempty, "type_counts": actual_type_counts, "csv_rows": actual_count})
    finally:
        formulas.close()
        cached.close()
    return {"ok": all(item["ok"] for item in checks), "sheets": checks}
